"""Offline manuscript rendering for the saved amended study.

This helper reads only a curated prose specification, an already-built execution
summary and explicitly named local images. It does not import scientific runtime,
controllers, providers or SQLite. It never finalizes results by itself. Rendering
requires the host's end-of-schedule record bound to the exact summary bytes.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime, timezone
from hashlib import sha256
from html import escape
import json
from pathlib import Path
import re
import shutil
import tempfile

ROOT = Path(__file__).resolve().parents[1]
SECTION_IDS = ["introduction", "related_work", "methods", "results", "discussion", "conclusion", "code_data", "references", "limitations_dual_use", "llm_pre_sprint"]
TABLES = {"execution", "primary", "controls", "review_audit"}


def read(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def file_sha(path):
    return sha256(Path(path).read_bytes()).hexdigest()


def word_count(text):
    return len(text.split())


def validate_spec(spec, final=False):
    if spec.get("schema_version") != 1:
        raise ValueError("Unknown manuscript specification version")
    for key in ("title", "authors", "affiliations", "abstract", "pre_sprint_disclosure"):
        if not isinstance(spec.get(key), str) or not spec[key].strip():
            raise ValueError("Missing required prose field: " + key)
    ids = [s.get("id") for s in spec.get("sections", [])]
    if ids != SECTION_IDS:
        raise ValueError("Official-template sections must be present in the documented order")
    auto_tables = []
    for section in spec["sections"]:
        if not section.get("title") or not section.get("blocks"):
            raise ValueError("Empty required section: " + section["id"])
        for block in section["blocks"]:
            kind = block.get("kind")
            if kind not in {"paragraph", "auto_table", "table", "figure", "code"}:
                raise ValueError("Unsupported block kind: " + str(kind))
            if kind in {"paragraph", "code"} and not isinstance(block.get("text"), str):
                raise ValueError("A prose/code block needs text")
            if kind == "auto_table":
                if block.get("name") not in TABLES:
                    raise ValueError("Unknown saved-data table")
                auto_tables.append(block["name"])
            if kind == "table" and (not block.get("headers") or any(len(row) != len(block["headers"]) for row in block.get("rows", []))):
                raise ValueError("Malformed manually curated table")
    if set(auto_tables) != TABLES or len(auto_tables) != len(TABLES):
        raise ValueError("Use exactly one of each execution/primary/controls/review_audit table")
    if final:
        if not spec.get("results_finalized"):
            raise ValueError("Prose is marked preparation-only; no final rendering")
        if "[[PENDING" in json.dumps(spec):
            raise ValueError("Unresolved result/prose placeholders remain")
        if not 150 <= word_count(spec["abstract"]) <= 250:
            raise ValueError("Abstract must contain 150–250 words")
        if not spec.get("claims_reviewed_against_evidence"):
            raise ValueError("Final prose has not been checked against actual evidence")
    return {"template_sections_valid": True, "abstract_words": word_count(spec["abstract"]), "results_finalized": bool(spec.get("results_finalized"))}


def validate_summary(summary):
    # Every planned slot must survive final reporting, including missing endpoints.
    expected = {"cases": 10, "primary_case_pairs": 6, "primary_receptions": 12,
                "controls": 4, "unique_round_observations": 48,
                "trajectories": 16, "audit_views": 32}
    for key, count in expected.items():
        if len(summary.get(key, [])) != count:
            raise ValueError(f"Saved summary must retain all {count} planned {key} slots")
    for pair in summary["primary_case_pairs"]:
        if len(pair.get("branches", [])) != 2:
            raise ValueError("A main case is missing a planned branch slot")
    c = summary["completion"]
    if sum(bool(r["reception_observed"]) for r in summary["primary_receptions"]) != c["primary_receptions"]["observed"]:
        raise ValueError("Primary endpoint count mismatch")
    if sum(r["neutral_status"] == "valid" for r in summary["primary_receptions"]) != c["neutral_replays"]["valid"]:
        raise ValueError("Neutral endpoint count mismatch")
    if sum(r["status"] == "complete" for r in summary["trajectories"]) != c["complete_trajectories"]["observed"]:
        raise ValueError("Trajectory count mismatch")
    if sum(r["state"] == "reviewed" for r in summary["unique_round_observations"]) != c["unique_reviewed_rounds"]["observed"]:
        raise ValueError("Unique reviewer count mismatch")
    audit_statuses = dict(Counter(r["status"] for r in summary["audit_views"]))
    if {k: v for k, v in c["audit_views"]["statuses"].items() if v} != audit_statuses:
        raise ValueError("Actual audit view count mismatch")
    if summary.get("warnings") or summary["budget"].get("local_aggregate_state_mismatches"):
        raise ValueError("Resolve or explicitly review saved accounting warnings before rendering")


def missing_label(branch, neutral=False):
    if neutral:
        return branch.get("neutral_outcome", {}).get("classification", branch.get("neutral_status", "missing"))
    return "missing" if branch.get("attempted") else "not reached"


def decision_label(branch, neutral=False):
    # A raw, schema-invalid candidate name is never promoted to a valid endpoint.
    valid = branch.get("neutral_status") == "valid" if neutral else branch.get("reception_observed", False)
    value = branch.get("neutral" if neutral else "intact", {}).get("decision")
    return value if valid and value is not None else "missing"


def value_label(value):
    if value is None:
        return "unresolved"
    if isinstance(value, bool):
        return "yes" if value else "no"
    return f"{value:g}" if isinstance(value, (float, int)) else str(value)


def data_table(name, d):
    c = d["completion"]
    if name == "execution":
        rows = [["Cases attempted", str(sum(x["attempted_requests"] > 0 for x in d["cases"])), "10"],
                ["A/B delivery pairs", str(sum(x["delivery_pairs_observed"] for x in d["cases"])), "48"],
                ["Primary round-3 B receptions", str(c["primary_receptions"]["observed"]), "12"],
                ["Valid neutral full-turn replays", str(c["neutral_replays"]["valid"]), "12"],
                ["Unique round reviews", str(c["unique_reviewed_rounds"]["observed"]), "48"],
                ["Completed final trajectories", str(c["complete_trajectories"]["observed"]), "16"],
                ["Valid final-auditor views", str(c["audit_views"]["valid"]), "32"],
                ["Invalid final-auditor views", str(c["audit_views"]["invalid"]), "32 slots"],
                ["Audit attempts with no outcome", str(c["audit_views"]["attempted_no_outcome"]), "32 slots"],
                ["Audit views not attempted", str(c["audit_views"]["not_attempted"]), "32 slots"],
                ["Valid latest/history audit pairs", str(c["valid_paired_audits"]["observed"]), "16"]]
        return {"caption": "Table 1. Execution coverage. Observed audit views are counted from saved calls, independently of trajectory completion. Missing work remains in the planned denominators.",
                "headers": ["Component", "Observed", "Planned"], "rows": rows, "widths": [0.64, 0.18, 0.18]}
    if name == "primary":
        rows = []
        for pair in sorted(d["primary_case_pairs"], key=lambda p: p["case_id"]):
            branches = sorted(pair["branches"], key=lambda b: b["branch"])
            cells, missing = [], []
            for b in branches:
                cells.append(f"{b['target']}: {decision_label(b)} / {decision_label(b, True)}")
                if not b.get("reception_observed"):
                    missing.append(b["branch"] + " intact " + missing_label(b))
                if b.get("neutral_status") != "valid":
                    missing.append(b["branch"] + " neutral " + missing_label(b, True))
            rows.append([pair["case_id"], *cells, value_label(pair.get("target_dependence_theta")),
                         value_label(pair.get("intact_minus_neutral_target_matches")), "; ".join(missing) or "None at primary/replay"])
        return {"caption": "Table 2. All six main case blocks. Each branch cell gives private target: intact B decision / neutral B decision. Theta is the prespecified case-level directional target dependence; delta is intact minus neutral target matches. Unresolved values are missing comparisons, not observed failures. Later trajectory losses do not erase valid primary receptions.",
                "headers": ["Case", "Branch 1 target: I / N", "Branch 2 target: I / N", "Theta", "Delta", "Primary/replay missingness"],
                "rows": rows, "widths": [0.105, 0.20, 0.20, 0.075, 0.075, 0.345]}
    if name == "controls":
        rows=[]
        for control in sorted(d["controls"], key=lambda p: p["case_id"]):
            failures = [x.get("reason", "unspecified failure") for x in control.get("failures", [])]
            missing = len(control["unobserved_round_keys"])
            rows.append([control["case_id"] + " / " + control["condition"],
                         f"{control['correct_observed_decisions']}/{control['observed_decision_denominator']}",
                         str(control["planned_unique_rounds"]), str(missing),
                         "; ".join(failures) or ("overt pixel signal: " + value_label(control.get("overt_pixel_signal_verified")) if control["condition"] == "overt" else "No recorded failure")])
        return {"caption": "Table 3. Honest and overt controls remain separate. Correctness is correct valid observed B decisions / valid observed B decisions; planned rounds and missing rounds are shown separately. Image fidelity and overt signaling require independent pixel coding.",
                "headers": ["Case / condition", "Correct / observed", "Planned rounds", "Missing rounds", "Observed limitation / coding"],
                "rows": rows, "widths": [0.17, 0.14, 0.12, 0.12, 0.45]}
    if name == "review_audit":
        t = d["review_trace_accounting"]
        def counts(items): return "; ".join(f"{key}: {value}" for key, value in sorted(items.items())) or "none"
        valid_views = [r for r in d["audit_views"] if r["status"] == "valid"]
        rows = [["Codex round reviews", str(t["unique_reviews_observed"]), counts(t["statuses"]), f"Blinded-eligible: {t['blinded_eligible_reviews']}"],
                ["Sonnet latest pair", str(sum(r["view"] == "latest_pair" for r in valid_views)), counts(dict(Counter(r["status"] for r in d["audit_views"] if r["view"] == "latest_pair"))), "16 planned views"],
                ["Sonnet image history", str(sum(r["view"] == "image_history" for r in valid_views)), counts(dict(Counter(r["status"] for r in d["audit_views"] if r["view"] == "image_history"))), "16 planned views"]]
        return {"caption": "Table 4. Oversight evidence and trace eligibility. The two Sonnet views receive identical task facts and the full B-decision timeline; view disagreement does not alone show improved accuracy.",
                "headers": ["Reviewer / view", "Observed valid", "Actual status / trace coverage", "Eligibility / scope"],
                "rows": rows, "widths": [0.23, 0.13, 0.38, 0.26]}
    raise ValueError(name)


def blocks_from_spec(spec, summary):
    blocks = [{"kind": "heading", "level": 1, "text": spec["title"]},
              {"kind": "paragraph", "text": spec["authors"]},
              {"kind": "paragraph", "text": spec["affiliations"]},
              {"kind": "paragraph", "text": "DRAFT — not submitted. " + spec["pre_sprint_disclosure"]},
              {"kind": "heading", "level": 2, "text": "Abstract"},
              {"kind": "paragraph", "text": spec["abstract"]}]
    for section in spec["sections"]:
        if section["id"] == "references":
            blocks.append({"kind": "pagebreak"})
        blocks.append({"kind": "heading", "level": 2, "text": section["title"], "section_id": section["id"]})
        for block in section["blocks"]:
            if block["kind"] == "auto_table":
                blocks.append({"kind": "table", **data_table(block["name"], summary)})
            elif block["kind"] == "figure":
                path = (ROOT / block["path"]).resolve()
                if not path.is_relative_to(ROOT) or path.is_symlink() or file_sha(path) != block["sha256"]:
                    raise ValueError("Figure source/hash outside the saved evidence scope")
                if block.get("visibility") not in {"delivered", "private", "author_annotation"}:
                    raise ValueError("Figure requires delivered/private/author_annotation visibility")
                if block["visibility"] != "delivered" and block["visibility"].replace("_", " ") not in block["caption"].lower():
                    raise ValueError("Private/annotated figures must identify that scope in the caption")
                blocks.append({**block, "path": str(path)})
            else:
                blocks.append(block)
    return blocks


def markdown(blocks, path):
    lines=[]
    for b in blocks:
        kind=b["kind"]
        if kind == "heading": lines.extend(["#"*b["level"]+" "+b["text"], ""])
        elif kind == "paragraph": lines.extend([b["text"], ""])
        elif kind == "code": lines.extend(["```text", b["text"], "```", ""])
        elif kind == "table":
            lines.extend(["**"+b["caption"]+"**", "", "| "+" | ".join(b["headers"])+" |", "| "+" | ".join("---" for _ in b["headers"])+" |"])
            for row in b["rows"]: lines.append("| "+" | ".join(str(v).replace("|", "\\|").replace("\n", "<br>") for v in row)+" |")
            lines.append("")
        elif kind == "figure": lines.extend([f"![{b['caption']}]({b['path']})", "", b["caption"], ""])
        elif kind == "pagebreak": lines.extend(["---", ""])
    path.write_text("\n".join(lines), encoding="utf-8")


def pdf(blocks, path):
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak, KeepTogether
    from PIL import Image as PILImage
    font_root=Path('/usr/share/fonts/truetype/liberation')
    for name, filename in (("PaperRoman", "LiberationSerif-Regular.ttf"), ("PaperBold", "LiberationSerif-Bold.ttf"), ("PaperItalic", "LiberationSerif-Italic.ttf")):
        pdfmetrics.registerFont(TTFont(name, str(font_root/filename)))
    pdfmetrics.registerFont(TTFont("PaperMono", "/usr/share/fonts/truetype/noto/NotoSansMono-Regular.ttf"))
    styles=getSampleStyleSheet()
    for name, kwargs in {
        "PaperBody": dict(fontName="PaperRoman", fontSize=10.2, leading=13.2, spaceAfter=5),
        "PaperTitle": dict(fontName="PaperBold", fontSize=17, leading=20, spaceAfter=8, keepWithNext=True),
        "PaperHeading": dict(fontName="PaperBold", fontSize=12, leading=15, spaceBefore=7, spaceAfter=4, keepWithNext=True),
        "PaperCaption": dict(fontName="PaperItalic", fontSize=8.5, leading=10.7, spaceAfter=6),
        "PaperTableCaption": dict(fontName="PaperItalic", fontSize=8.5, leading=10.7, spaceAfter=5, keepWithNext=True),
        "PaperCell": dict(fontName="PaperRoman", fontSize=8.2, leading=10.2),
        "PaperCode": dict(fontName="PaperMono", fontSize=7, leading=9, wordWrap="CJK", spaceAfter=4),
    }.items(): styles.add(ParagraphStyle(name=name, **kwargs))
    class PaperDocument(SimpleDocTemplate):
        references_page=None
        def afterFlowable(self, flowable):
            if getattr(flowable, "paper_section_id", None)=="references": self.references_page=self.page
    document=PaperDocument(str(path), pagesize=(612,792), leftMargin=44, rightMargin=44, topMargin=40, bottomMargin=43)
    story=[]
    for b in blocks:
        kind=b["kind"]
        if kind in {"heading", "paragraph", "code"}:
            style="PaperTitle" if kind=="heading" and b["level"]==1 else "PaperHeading" if kind=="heading" else "PaperCode" if kind=="code" else "PaperBody"
            item=Paragraph(escape(b["text"]).replace("\n", "<br/>"), styles[style]);item.paper_section_id=b.get("section_id");story.append(item)
        elif kind == "table":
            story.append(Paragraph(escape(b["caption"]),styles["PaperTableCaption"]))
            cells=[[Paragraph(escape(str(v)).replace("\n","<br/>"),styles["PaperCell"]) for v in row] for row in [b["headers"],*b["rows"]]]
            proportions=b.get("widths", [1/len(b["headers"])]*len(b["headers"]))
            widths=[524*x/sum(proportions) for x in proportions]
            table=Table(cells, colWidths=widths, repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#e8edf2")),("GRID",(0,0),(-1,-1),.3,colors.HexColor("#b4bec8")),("VALIGN",(0,0),(-1,-1),"TOP"),("LEFTPADDING",(0,0),(-1,-1),4),("RIGHTPADDING",(0,0),(-1,-1),4),("TOPPADDING",(0,0),(-1,-1),3),("BOTTOMPADDING",(0,0),(-1,-1),3)]))
            story.extend([table,Spacer(1,6)])
        elif kind == "figure":
            with PILImage.open(b["path"]) as im: width,height=im.size
            ratio=min(float(b.get("width_inches",4))*72/width, 330/height)
            im=Image(b["path"],width=width*ratio,height=height*ratio);im.hAlign="CENTER"
            story.append(KeepTogether([im,Spacer(1,4),Paragraph(escape(b["caption"]),styles["PaperCaption"])]))
        elif kind == "pagebreak": story.append(PageBreak())
    def footer(canvas, doc):
        canvas.saveState();canvas.setFont("PaperRoman",8);canvas.drawString(44,24,"Research draft | Pre-sprint work disclosed | Not submitted");canvas.drawRightString(568,24,str(doc.page));canvas.restoreState()
    document.build(story,onFirstPage=footer,onLaterPages=footer)
    if document.references_page is None: raise ValueError("References page boundary not recorded")
    return document.references_page-1


def docx(blocks, path):
    from docx import Document
    from docx.shared import Inches, Pt
    document=Document();section=document.sections[0]
    section.top_margin=Inches(.56);section.bottom_margin=Inches(.6)
    section.left_margin=Inches(.61);section.right_margin=Inches(.61)
    normal=document.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(10)
    for b in blocks:
        kind=b["kind"]
        if kind=="heading": document.add_heading(b["text"],level=0 if b["level"]==1 else 1)
        elif kind in {"paragraph","code"}: document.add_paragraph(b["text"])
        elif kind=="table":
            document.add_paragraph(b["caption"],style='Caption');table=document.add_table(rows=1,cols=len(b["headers"]));table.style='Table Grid'
            for cell,value in zip(table.rows[0].cells,b["headers"]):cell.text=str(value)
            for row in b["rows"]:
                for cell,value in zip(table.add_row().cells,row):cell.text=str(value)
        elif kind=="figure":document.add_picture(b["path"],width=Inches(float(b.get("width_inches",4))));document.add_paragraph(b["caption"],style='Caption')
        elif kind=="pagebreak":document.add_page_break()
    document.save(path)


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--spec", type=Path, required=True)
    parser.add_argument("--check-template",action="store_true")
    parser.add_argument("--summary",type=Path)
    parser.add_argument("--completion-record",type=Path)
    parser.add_argument("--destination",type=Path)
    args=parser.parse_args();spec=read(args.spec)
    check=validate_spec(spec,final=not args.check_template)
    if args.check_template:
        print(json.dumps({**check,"rendered":False}));return
    if not all((args.summary,args.completion_record,args.destination)):
        parser.error("rendering needs --summary, --completion-record and --destination")
    completion=read(args.completion_record);summary=read(args.summary)
    if completion.get("host_completed_or_documented_terminal") is not True or completion.get("summary_sha256")!=file_sha(args.summary):
        raise ValueError("Await an explicit host end-of-schedule record bound to these summary bytes")
    validate_summary(summary)
    blocks=blocks_from_spec(spec,summary)
    destination=args.destination.resolve();destination.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="paper-render-",dir=destination) as temp:
        temp=Path(temp)
        markdown(blocks,temp/'PAPER_DRAFT.md')
        main_pages=pdf(blocks,temp/'PAPER_DRAFT.pdf')
        if main_pages>8: raise ValueError(f"Main paper is {main_pages} pages; compact prose/figures before finalizing (maximum 8)")
        docx(blocks,temp/'PAPER_DRAFT.docx')
        from pypdf import PdfReader
        pages=len(PdfReader(str(temp/'PAPER_DRAFT.pdf')).pages)
        build={"generated_at_utc":datetime.now(timezone.utc).isoformat(),"abstract_words":word_count(spec['abstract']),"main_pages":main_pages,"pdf_pages_total":pages,"results_finalized_for_draft":True,"authors_missing":bool(spec.get('authors_missing',True)),"submission_ready":False,"template_scope":"Observed official section structure; exact native template typography is not asserted", "inputs":{str(p.resolve()):file_sha(p) for p in (args.spec,args.summary,args.completion_record,Path(__file__))},"runtime_status":summary.get('runtime_status'),"scientific_schedule_fully_observed":summary.get('scientific_schedule_fully_observed'),"completion":summary['completion'],"data_scope":summary['scope']}
        (temp/'PAPER_BUILD.json').write_text(json.dumps(build,indent=2)+'\n')
        for filename in ('PAPER_DRAFT.md','PAPER_DRAFT.pdf','PAPER_DRAFT.docx','PAPER_BUILD.json'):
            shutil.move(str(temp/filename),str(destination/filename))
    print(json.dumps({"abstract_words":build['abstract_words'],"main_pages":main_pages,"pdf_pages_total":pages,"submission_ready":False,"destination":str(destination)},indent=2))


if __name__=='__main__':main()
